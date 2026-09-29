# -*- coding: utf-8 -*-
"""
Génère docs/tco-verification.xlsx : classeur d'audit des 7 cas de référence
TCO — méthodologie v2.2 (docs/tco-methodologie.md).

SOURCE UNIQUE : toutes les hypothèses communes ET les entrées des cas sont
lues dans docs/tco-cas-de-reference.json, lui-même GÉNÉRÉ par
scripts/reference-cases/generate.mjs depuis cas-definitions.mjs — dont les
paramètres sont testés IDENTIQUES à src/lib/tco/assumptions.ts
(defaults.test.ts). Rien n'est recopié à la main ici.

Principe : une feuille « Hypothèses » (cellules nommées), puis une feuille
par cas où TOUTES les cellules de calcul sont des FORMULES Excel. Seules
les entrées du cas sont des constantes ; les « sorties attendues » du
contre-calculateur indépendant sont affichées EN REGARD des formules pour
comparaison, jamais utilisées dans un calcul.

Méthode v2.2 : taxes SYMÉTRIQUES (part non récupérable sur acquisition,
infra, énergie, entretien, opex — §3.1) ; VR géométrique planchée partout,
reprise au re-remplacement comprise (§10.2) ; infrastructure ré-investie en
fin de durée de vie si l'horizon la dépasse, VR linéaire du dernier
équipement (§3.5).

Usage : python3 scripts/tco-verification-xlsx.py
"""
import json
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

RACINE = os.path.join(os.path.dirname(__file__), "..")
OUT = os.path.join(RACINE, "docs", "tco-verification.xlsx")
REF = json.load(open(os.path.join(RACINE, "docs", "tco-cas-de-reference.json"), encoding="utf-8"))
P = REF["parametresCommuns"]

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
FILL_IN = PatternFill("solid", fgColor="FFF2CC")    # entrées du cas (constantes)
FILL_CALC = PatternFill("solid", fgColor="DDEBF7")  # formules
FILL_EXP = PatternFill("solid", fgColor="E2EFDA")   # sorties attendues (contre-calcul)

# ---------------------------------------------------------------------------
# Feuille Hypothèses : (nom défini, libellé, valeur ou formule, unité) —
# VALEURS LUES dans le JSON généré, jamais recopiées.
# ---------------------------------------------------------------------------
FE = P["facteursEmission"]
HYP = [
    ("taux_r",             "Taux d'actualisation nominal r", P["tauxActualisationNominal"], "décimal"),
    ("g_diesel",           "Inflation diesel", P["inflations"]["diesel"], "décimal"),
    ("g_elec",             "Inflation électricité", P["inflations"]["electricite"], "décimal"),
    ("g_h2",               "Inflation hydrogène", P["inflations"]["hydrogene"], "décimal"),
    ("g_entretien",        "Inflation entretien", P["inflations"]["entretien"], "décimal"),
    ("g_general",          "Inflation générale", P["inflations"]["generale"], "décimal"),
    ("prix_diesel",        "Prix diesel année 0 (AVANT TPS/TVQ, accises comprises — moyenne 12 mois StatCan)",
     P["prixAnnee0"]["dieselParL"], "$/L"),
    ("prix_elec",          "Prix électricité effectif au compteur, année 0 (avant taxes)",
     P["prixAnnee0"]["electriciteEffectiveParKwh"], "$/kWh"),
    ("prix_h2",            "Prix H2 livré, année 0 (avant taxes)", P["prixAnnee0"]["h2LivreParKg"], "$/kg"),
    ("rendement",          "Rendement de recharge (kWh véhicule / kWh compteur)", P["rendementRecharge"], "ratio"),
    ("maj_annualisee",     "Majoration hivernale ANNUALISÉE (majoration x part des km d'hiver)",
     P["majorationHivernaleAnnualisee"], "ratio"),
    ("fact_hiver",         "Facteur hivernal = 1 + majoration annualisée", "=1+maj_annualisee", "ratio"),
    ("taux_taxes",         "Taxes non récupérables (municipalité : TPS remboursée 100 %, TVQ à 50 % — §3.1 v2.2, "
                           "appliquées SYMÉTRIQUEMENT à tous les postes)", P["tauxTaxesNonRecuperables"], "décimal"),
    ("d_diesel",           "Dépréciation annuelle diesel", P["depreciationAnnuelle"]["diesel"], "décimal"),
    ("d_bev",              "Dépréciation annuelle BEV", P["depreciationAnnuelle"]["BEV"], "décimal"),
    ("d_fcev",             "Dépréciation annuelle FCEV", P["depreciationAnnuelle"]["FCEV"], "décimal"),
    ("plancher_vr",        "Plancher de valeur résiduelle", P["plancherResiduel"], "ratio"),
    ("infra_entretien_pct", "Entretien infra annuel (part du capex avant taxes)",
     P["infra"]["entretienAnnuelPctCapex"], "ratio"),
    ("duree_vie_infra",    "Durée de vie de l'infrastructure", P["infra"]["dureeVieAns"], "années"),
    ("fe_leger",           "FE diesel TTW — véhicules légers/camionnettes", FE["dieselTtwLegersKgParL"], "kg CO2e/L"),
    ("fe_lourd",           "FE diesel TTW — camions moyens/lourds, autobus", FE["dieselTtwLourdsKgParL"], "kg CO2e/L"),
    ("ratio_wtw",          "WTW diesel = TTW x ratio", FE["ratioWtwDiesel"], "ratio"),
    ("fe_elec",            "FE électricité QC (WTW, au compteur)", FE["electriciteGParKwh"], "g CO2e/kWh"),
    ("fe_h2",              "FE H2 électrolyse QC (WTW)", FE["h2KgParKg"], "kg CO2e/kg"),
]


def build_hypotheses(wb):
    ws = wb.active
    ws.title = "Hypothèses"
    ws["A1"] = ("Hypothèses communes — scénario Central, CAD courants, année de référence "
                f"{P['anneeReference']}, municipalité")
    ws["A1"].font = TITLE
    ws["A2"] = ("GÉNÉRÉES depuis docs/tco-cas-de-reference.json (lui-même généré depuis "
                "scripts/reference-cases/cas-definitions.mjs, testé identique à src/lib/tco/assumptions.ts) — "
                "ne pas éditer à la main.")
    ws.append([])
    ws.append(["Nom", "Description", "Valeur", "Unité"])
    for c in "ABCD":
        ws[f"{c}4"].font = BOLD
    row = 5
    for nom, desc, val, unite in HYP:
        ws.cell(row=row, column=1, value=nom)
        ws.cell(row=row, column=2, value=desc)
        cell = ws.cell(row=row, column=3, value=val)
        cell.fill = FILL_CALC if isinstance(val, str) and val.startswith("=") else FILL_IN
        ws.cell(row=row, column=4, value=unite)
        wb.defined_names[nom] = DefinedName(nom, attr_text=f"'Hypothèses'!$C${row}")
        row += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 84
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 14
    return ws


# ---------------------------------------------------------------------------
# Feuilles des cas à UN véhicule (1 à 5 et 7) — spécifications extraites
# du JSON des cas de référence.
# ---------------------------------------------------------------------------
COLS = ["n", "Acquisition alt", "Subventions", "Énergie alt", "Entretien alt",
        "Opex infra", "Résiduels alt", "Flux net alt", "Acquisition réf",
        "Énergie réf", "Entretien réf", "Résiduels réf", "Flux net réf",
        "Économie nominale", "Facteur actualisation", "Flux act. alt",
        "Flux act. réf", "Cumul économie nominale", "Cumul économie actualisée",
        "km actualisés", "n si cumul nominal >= 0", "n si cumul actualisé >= 0",
        "TTW réf (t)", "WTW réf (t)", "WTW alt (t)"]


def spec_depuis_json(cas):
    """Cas à un seul véhicule : entrées du JSON → spécification de feuille."""
    e = cas["entrees"]
    assert len(e["vehicules"]) == 1, f"cas {cas['id']} n'est pas mono-véhicule"
    v = e["vehicules"][0]
    subs = v.get("subventionsAlternative") or []
    assert len(subs) <= 1, f"cas {cas['id']} : plus d'une subvention, feuille non prévue"
    sites = e.get("sitesInfra") or []
    return dict(
        id=cas["id"], titre=cas["titre"],
        H=e["parametres"]["horizonAns"],
        km=v["kmParAn"],
        ref_prix=v["reference"]["prixAvantTaxes"], ref_conso=v["reference"]["consommationPar100km"],
        ref_cpk=v["reference"]["entretienParKm"],
        techno=v["alternative"]["technologie"],
        alt_prix=v["alternative"]["prixAvantTaxes"], alt_conso=v["alternative"]["consommationPar100km"],
        alt_cpk=v["alternative"]["entretienParKm"],
        sub=(subs[0]["montant"] if subs else 0), sub_annee=(subs[0]["annee"] if subs else 1),
        infra=(sites[0]["capexAvantTaxes"] if sites else 0),
        fe=("fe_leger" if v["classeEmissionDiesel"] == "legers" else "fe_lourd"),
        sorties=cas["sorties"],
    )


def build_case_simple(wb, spec):
    ws = wb.create_sheet(f"Cas {spec['id']}")
    ws["A1"] = f"Cas {spec['id']} — {spec['titre']}"
    ws["A1"].font = TITLE

    techno = spec["techno"]  # 'BEV' | 'FCEV'
    inputs = [
        ("Horizon H (années)", spec["H"], "ans"),
        ("Kilométrage annuel", spec["km"], "km/an"),
        ("Prix réf (diesel) avant taxes", spec["ref_prix"], "$"),
        ("Conso réf", spec["ref_conso"], "L/100 km"),
        ("Entretien réf", spec["ref_cpk"], "$/km"),
        ("Prix alt avant taxes", spec["alt_prix"], "$"),
        ("Conso alt (nominale tempérée)", spec["alt_conso"], "kWh/100 km" if techno == "BEV" else "kg H2/100 km"),
        ("Entretien alt", spec["alt_cpk"], "$/km"),
        ("Subvention (montant)", spec["sub"], "$"),
        ("Subvention (année de versement)", spec["sub_annee"], "année"),
        ("Infra — capex avant taxes", spec["infra"], "$"),
        ("FE diesel TTW (choix catégorie)", f"={spec['fe']}", "kg CO2e/L"),
    ]
    ws["A3"] = "Entrées du cas (constantes, depuis docs/tco-cas-de-reference.json)"
    ws["A3"].font = BOLD
    r = 4
    for lab, val, unite in inputs:
        ws.cell(row=r, column=1, value=lab)
        c = ws.cell(row=r, column=2, value=val)
        c.fill = FILL_CALC if isinstance(val, str) and str(val).startswith("=") else FILL_IN
        ws.cell(row=r, column=3, value=unite)
        r += 1
    # B4..B15
    H_, KM, PREF, CREF, EREF = "$B$4", "$B$5", "$B$6", "$B$7", "$B$8"
    PALT, CALT, EALT, SUB, SUBAN, INFRA, FE_ = "$B$9", "$B$10", "$B$11", "$B$12", "$B$13", "$B$14", "$B$15"

    d_alt = {"BEV": "d_bev", "FCEV": "d_fcev"}[techno]
    conso_compteur = (f"={CALT}*fact_hiver/rendement" if techno == "BEV"
                      else f"={CALT}*fact_hiver")
    prix_energie_alt = "prix_elec" if techno == "BEV" else "prix_h2"
    g_energie_alt = "g_elec" if techno == "BEV" else "g_h2"

    derived = [
        ("Prix réf avec taxes non récupérables", f"={PREF}*(1+taux_taxes)"),
        ("Prix alt avec taxes non récupérables", f"={PALT}*(1+taux_taxes)"),
        ("Infra avec taxes non récupérables", f"={INFRA}*(1+taux_taxes)"),
        ("Conso alt facturée (compteur ou pompe) /100 km", conso_compteur),
        ("VR réf fin d'horizon (avant taxes, nominale, géométrique planchée)",
         f"=MAX({PREF}*(1-d_diesel)^{H_},plancher_vr*{PREF})"),
        ("VR alt fin d'horizon (avant taxes, nominale, géométrique planchée)",
         f"=MAX({PALT}*(1-{d_alt})^{H_},plancher_vr*{PALT})"),
        # §3.5 v2.2 : si H > durée de vie de l'infra, ré-achat en fin de
        # vie (capex indexé à l'inflation générale) et VR LINÉAIRE du
        # dernier équipement en fin d'horizon ; sinon VR linéaire de
        # l'équipement initial.
        ("Année de ré-achat de l'infra (vide si aucun)",
         f'=IF(AND({INFRA}>0,{H_}>duree_vie_infra),duree_vie_infra,"")'),
        ("VR infra fin d'horizon (linéaire, dernier équipement, capex indexé)",
         f"=IF({INFRA}=0,0,IF({H_}>duree_vie_infra,"
         f"{INFRA}*(1+g_general)^duree_vie_infra*(duree_vie_infra-({H_}-duree_vie_infra))/duree_vie_infra,"
         f"{INFRA}*(duree_vie_infra-{H_})/duree_vie_infra))"),
    ]
    ws.cell(row=17, column=1, value="Valeurs dérivées (formules)").font = BOLD
    r = 18
    for lab, f in derived:
        ws.cell(row=r, column=1, value=lab)
        c = ws.cell(row=r, column=2, value=f); c.fill = FILL_CALC
        r += 1
    PREF_TTC, PALT_TTC, INFRA_TTC = "$B$18", "$B$19", "$B$20"
    CONSO_C, VRREF, VRALT = "$B$21", "$B$22", "$B$23"
    AN_REACHAT, VRINFRA = "$B$24", "$B$25"

    # tableau annuel — TAXES SYMÉTRIQUES (§3.1 v2.2) : énergie, entretien
    # et opex portent (1+taux_taxes) comme les acquisitions.
    hdr = 28
    for j, name in enumerate(COLS, start=1):
        c = ws.cell(row=hdr, column=j, value=name); c.font = BOLD
    first = hdr + 1
    H = spec["H"]
    for n in range(H + 1):
        row = first + n
        A = f"A{row}"
        ws.cell(row=row, column=1, value=(0 if n == 0 else f"=A{row-1}+1"))
        formulas = [
            # B acq alt : an 0 = véhicule + infra ; ré-achat de l'infra en
            # fin de vie (capex indexé, taxes) si l'horizon la dépasse
            f"=IF({A}=0,{PALT_TTC}+{INFRA_TTC},"
            f"IF({A}={AN_REACHAT},{INFRA}*(1+g_general)^{A}*(1+taux_taxes),0))",
            f"=IF({A}={SUBAN},{SUB},0)",                                          # C subventions
            f"=IF({A}=0,0,{KM}*{CONSO_C}/100*{prix_energie_alt}*(1+{g_energie_alt})^{A}*(1+taux_taxes))",  # D
            f"=IF({A}=0,0,{KM}*{EALT}*(1+g_entretien)^{A}*(1+taux_taxes))",       # E entretien alt
            f"=IF({A}=0,0,infra_entretien_pct*{INFRA}*(1+g_entretien)^{A}*(1+taux_taxes))",  # F opex infra
            f"=IF({A}={H_},{VRALT}+{VRINFRA},0)",                                 # G résiduels alt
            f"=B{row}-C{row}+D{row}+E{row}+F{row}-G{row}",                        # H flux alt
            f"=IF({A}=0,{PREF_TTC},0)",                                           # I acq réf
            f"=IF({A}=0,0,{KM}*{CREF}/100*prix_diesel*(1+g_diesel)^{A}*(1+taux_taxes))",  # J énergie réf
            f"=IF({A}=0,0,{KM}*{EREF}*(1+g_entretien)^{A}*(1+taux_taxes))",       # K entretien réf
            f"=IF({A}={H_},{VRREF},0)",                                           # L résiduels réf
            f"=I{row}+J{row}+K{row}-L{row}",                                      # M flux réf
            f"=M{row}-H{row}",                                                    # N économie
            f"=1/(1+taux_r)^{A}",                                                 # O facteur
            f"=H{row}*O{row}",                                                    # P flux act alt
            f"=M{row}*O{row}",                                                    # Q flux act réf
            (f"=N{row}" if n == 0 else f"=R{row-1}+N{row}"),                      # R cumul nominal
            (f"=N{row}*O{row}" if n == 0 else f"=S{row-1}+N{row}*O{row}"),        # S cumul actualisé
            f"=IF({A}=0,0,{KM}*O{row})",                                          # T km actualisés
            f'=IF(R{row}>=0,{A},"")',                                             # U payback simple
            f'=IF(S{row}>=0,{A},"")',                                             # V payback actualisé
            f"=IF({A}=0,0,{KM}*{CREF}/100*{FE_}/1000)",                           # W TTW réf t
            f"=W{row}*ratio_wtw",                                                 # X WTW réf t
            (f"=IF({A}=0,0,{KM}*{CONSO_C}/100*fe_elec/1000000)" if techno == "BEV"
             else f"=IF({A}=0,0,{KM}*{CONSO_C}/100*fe_h2/1000)"),                 # Y WTW alt t
        ]
        for j, f in enumerate(formulas, start=2):
            c = ws.cell(row=row, column=j, value=f); c.fill = FILL_CALC
    last = first + H

    write_results(ws, first, last, res_row=last + 3, sorties=spec.get("sorties"))
    style_table(ws, hdr, last)
    return ws


def write_results(ws, first, last, res_row, sorties=None):
    ws.cell(row=res_row - 1, column=1, value="Résultats (formules)").font = BOLD
    if sorties is not None:
        c = ws.cell(row=res_row - 1, column=3,
                    value="Attendu (contre-calcul indépendant, JSON) — comparaison seulement")
        c.font = BOLD
    res = [
        ("TCO actualisé — alternative ($)", f"=SUM(P{first}:P{last})", "tcoActualiseAlt"),
        ("TCO actualisé — référence ($)", f"=SUM(Q{first}:Q{last})", "tcoActualiseRef"),
        ("VAN différentielle réf − alt ($)", f"=B{res_row+1}-B{res_row}", "vanDifferentielle"),
        ("km actualisés", f"=SUM(T{first}:T{last})", None),
        ("TCO/km actualisé — alternative ($/km)", f"=B{res_row}/B{res_row+3}", "tcoParKmAlt"),
        ("TCO/km actualisé — référence ($/km)", f"=B{res_row+1}/B{res_row+3}", None),
        ("Délai de récupération simple (années)",
         f'=IF(COUNT(U{first}:U{last})=0,"null : jamais atteint sur l\'horizon",MIN(U{first}:U{last}))',
         "paybackSimpleAns"),
        ("Délai de récupération actualisé (années)",
         f'=IF(COUNT(V{first}:V{last})=0,"null : jamais atteint sur l\'horizon",MIN(V{first}:V{last}))',
         "paybackActualiseAns"),
        ("CO2e évité cumulé TTW (t)", f"=SUM(W{first}:W{last})", "co2EviteTtwTonnes"),
        ("CO2e évité cumulé WTW (t)", f"=SUM(X{first}:X{last})-SUM(Y{first}:Y{last})", "co2EviteWtwTonnes"),
        ("Coût (+) ou gain (−) par tonne WTW évitée ($/t)",
         f"=(B{res_row}-B{res_row+1})/B{res_row+9}", "coutParTonneWtw"),
    ]
    for i, (lab, f, cle) in enumerate(res):
        ws.cell(row=res_row + i, column=1, value=lab).font = BOLD
        c = ws.cell(row=res_row + i, column=2, value=f); c.fill = FILL_CALC
        if sorties is not None and cle is not None and cle in sorties:
            attendu = sorties[cle]
            ce = ws.cell(row=res_row + i, column=3,
                         value=("null" if attendu is None else attendu))
            ce.fill = FILL_EXP


def style_table(ws, hdr, last):
    ws.column_dimensions["A"].width = 46
    for j in range(2, len(COLS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    for row in ws.iter_rows(min_row=hdr + 1, max_row=last, min_col=2, max_col=len(COLS)):
        for c in row:
            if c.column <= 20:  # colonnes monétaires/facteurs
                c.number_format = "#,##0.00"
            else:
                c.number_format = "0.000"


# ---------------------------------------------------------------------------
# Feuille cas 6 (mini-plan 5 véhicules, H=12, infra partagée,
# re-remplacement an 10 avec REPRISE GÉOMÉTRIQUE PLANCHÉE — §10.2 v2.0)
# ---------------------------------------------------------------------------
def build_case6(wb, cas):
    e = cas["entrees"]
    H = e["parametres"]["horizonAns"]
    vehs = {v["id"]: v for v in e["vehicules"]}
    vl, cam, cm = vehs["vl1"], vehs["camionnette"], vehs["cm1"]
    infra_capex = e["sitesInfra"][0]["capexAvantTaxes"]

    ws = wb.create_sheet("Cas 6")
    ws["A1"] = f"Cas 6 — {cas['titre']}"
    ws["A1"].font = TITLE

    def put(row, lab, val, unite=""):
        ws.cell(row=row, column=1, value=lab)
        c = ws.cell(row=row, column=2, value=val)
        c.fill = FILL_CALC if isinstance(val, str) and str(val).startswith("=") else FILL_IN
        if unite:
            ws.cell(row=row, column=3, value=unite)

    sub1 = lambda v: (v.get("subventionsAlternative") or [{}])
    ws["A3"] = "Entrées du cas (constantes, depuis docs/tco-cas-de-reference.json)"; ws["A3"].font = BOLD
    put(4, "Horizon H", H, "ans")
    ws["A5"] = "Véhicules légers BEV (x2)"; ws["A5"].font = BOLD
    put(6, "Nombre", 2); put(7, "km/an (par véhicule)", vl["kmParAn"])
    put(8, "Prix réf avant taxes", vl["reference"]["prixAvantTaxes"], "$")
    put(9, "Conso réf", vl["reference"]["consommationPar100km"], "L/100 km")
    put(10, "Entretien réf", vl["reference"]["entretienParKm"], "$/km")
    put(11, "Prix alt avant taxes", vl["alternative"]["prixAvantTaxes"], "$")
    put(12, "Conso alt", vl["alternative"]["consommationPar100km"], "kWh/100 km")
    put(13, "Entretien alt", vl["alternative"]["entretienParKm"], "$/km")
    put(14, "Subventions an 0 par véhicule", sub1(vl)[0].get("montant", 0), "$")
    put(15, "Année de re-remplacement (fin de vie)", vl["dureeVieAns"])
    ws["A17"] = "Camionnette BEV"; ws["A17"].font = BOLD
    put(18, "km/an", cam["kmParAn"])
    put(19, "Prix réf avant taxes", cam["reference"]["prixAvantTaxes"], "$")
    put(20, "Conso réf", cam["reference"]["consommationPar100km"], "L/100 km")
    put(21, "Entretien réf", cam["reference"]["entretienParKm"], "$/km")
    put(22, "Prix alt avant taxes", cam["alternative"]["prixAvantTaxes"], "$")
    put(23, "Conso alt", cam["alternative"]["consommationPar100km"], "kWh/100 km")
    put(24, "Entretien alt", cam["alternative"]["entretienParKm"], "$/km")
    put(25, "Subvention (Écocamionnage 2b)", sub1(cam)[0].get("montant", 0), "$")
    put(26, "Année de versement", sub1(cam)[0].get("annee", 1))
    ws["A28"] = "Camions moyens BEV (x2)"; ws["A28"].font = BOLD
    put(29, "Nombre", 2); put(30, "km/an (par véhicule)", cm["kmParAn"])
    put(31, "Prix réf avant taxes", cm["reference"]["prixAvantTaxes"], "$")
    put(32, "Conso réf", cm["reference"]["consommationPar100km"], "L/100 km")
    put(33, "Entretien réf", cm["reference"]["entretienParKm"], "$/km")
    put(34, "Prix alt avant taxes", cm["alternative"]["prixAvantTaxes"], "$")
    put(35, "Conso alt", cm["alternative"]["consommationPar100km"], "kWh/100 km")
    put(36, "Entretien alt", cm["alternative"]["entretienParKm"], "$/km")
    put(37, "Subvention par véhicule", sub1(cm)[0].get("montant", 0), "$")
    put(38, "Année de versement", sub1(cm)[0].get("annee", 1))
    ws["A40"] = "Infrastructure partagée (mise en service an 0)"; ws["A40"].font = BOLD
    put(41, "Capex infra avant taxes", infra_capex, "$")

    NV, KMV, PRV, CRV, ERV = "$B$6", "$B$7", "$B$8", "$B$9", "$B$10"
    PAV, CAV, EAV, SUBV, ANREPL = "$B$11", "$B$12", "$B$13", "$B$14", "$B$15"
    KMC, PRC, CRC, ERC = "$B$18", "$B$19", "$B$20", "$B$21"
    PAC, CAC, EAC, SUBC, ANC = "$B$22", "$B$23", "$B$24", "$B$25", "$B$26"
    NM, KMM, PRM, CRM, ERM = "$B$29", "$B$30", "$B$31", "$B$32", "$B$33"
    PAM, CAM, EAM, SUBM, ANM = "$B$34", "$B$35", "$B$36", "$B$37", "$B$38"
    INFRA, H_ = "$B$41", "$B$4"

    ws["A43"] = "Valeurs dérivées (formules)"; ws["A43"].font = BOLD
    derived = [
        ("kWh compteur/an — 1 véhicule léger", f"={KMV}*{CAV}/100*fact_hiver/rendement"),
        ("kWh compteur/an — camionnette", f"={KMC}*{CAC}/100*fact_hiver/rendement"),
        ("kWh compteur/an — 1 camion moyen", f"={KMM}*{CAM}/100*fact_hiver/rendement"),
        ("kWh compteur/an — total site (année 1)", f"={NV}*B44+B45+{NM}*B46"),
        ("Prix rachat VL alt an 10 (avant taxes, inflation générale)", f"={PAV}*(1+g_general)^{ANREPL}"),
        ("Prix rachat VL réf an 10 (avant taxes)", f"={PRV}*(1+g_general)^{ANREPL}"),
        # §10.2 v2.0 : reprise du véhicule remplacé à sa VR GÉOMÉTRIQUE
        # PLANCHÉE (même méthode qu'en fin d'horizon), sur le prix d'achat
        # de l'ancien véhicule (acheté an 0, non indexé).
        ("VR an 10 — VL alt ancien (géométrique planchée, âge 10)",
         f"=MAX((1-d_bev)^{ANREPL},plancher_vr)*{PAV}"),
        ("VR an 10 — VL réf ancien (géométrique planchée, âge 10)",
         f"=MAX((1-d_diesel)^{ANREPL},plancher_vr)*{PRV}"),
        ("VR fin 12 — VL alt nouveau (2 ans de possession)",
         f"=MAX(B48*(1-d_bev)^({H_}-{ANREPL}),plancher_vr*B48)"),
        ("VR fin 12 — VL réf nouveau", f"=MAX(B49*(1-d_diesel)^({H_}-{ANREPL}),plancher_vr*B49)"),
        ("VR fin 12 — camionnette alt", f"=MAX({PAC}*(1-d_bev)^{H_},plancher_vr*{PAC})"),
        ("VR fin 12 — camionnette réf", f"=MAX({PRC}*(1-d_diesel)^{H_},plancher_vr*{PRC})"),
        ("VR fin 12 — 1 camion moyen alt", f"=MAX({PAM}*(1-d_bev)^{H_},plancher_vr*{PAM})"),
        ("VR fin 12 — 1 camion moyen réf", f"=MAX({PRM}*(1-d_diesel)^{H_},plancher_vr*{PRM})"),
        ("VR fin 12 — infrastructure (linéaire, âge 12 sur 15)",
         f"={INFRA}*(duree_vie_infra-{H_})/duree_vie_infra"),
    ]
    r = 44
    for lab, f in derived:
        ws.cell(row=r, column=1, value=lab)
        c = ws.cell(row=r, column=2, value=f); c.fill = FILL_CALC
        r += 1
    KWH_TOT = "$B$47"
    PRACHAT_A, PRACHAT_R = "$B$48", "$B$49"
    VR10A, VR10R = "$B$50", "$B$51"
    VR12_VLA, VR12_VLR = "$B$52", "$B$53"
    VR12_CA, VR12_CR = "$B$54", "$B$55"
    VR12_MA, VR12_MR, VR12_I = "$B$56", "$B$57", "$B$58"

    # tableau annuel — TAXES SYMÉTRIQUES (§3.1 v2.2)
    hdr = 61
    for j, name in enumerate(COLS, start=1):
        c = ws.cell(row=hdr, column=j, value=name); c.font = BOLD
    first = hdr + 1
    for n in range(H + 1):
        row = first + n
        A = f"A{row}"
        ws.cell(row=row, column=1, value=(0 if n == 0 else f"=A{row-1}+1"))
        formulas = [
            # B acq alt : an 0 = 5 véhicules + infra ; an 10 = rachat 2 VL
            f"=IF({A}=0,({NV}*{PAV}+{PAC}+{NM}*{PAM}+{INFRA})*(1+taux_taxes),"
            f"IF({A}={ANREPL},{NV}*{PRACHAT_A}*(1+taux_taxes),0))",
            # C subventions : an 0 = 2 x VL ; an 1 = camionnette + 2 x CM
            f"=IF({A}=0,{NV}*{SUBV},IF({A}=1,{SUBC}+{NM}*{SUBM},0))",
            f"=IF({A}=0,0,{KWH_TOT}*prix_elec*(1+g_elec)^{A}*(1+taux_taxes))",     # D énergie alt
            f"=IF({A}=0,0,({NV}*{KMV}*{EAV}+{KMC}*{EAC}+{NM}*{KMM}*{EAM})*(1+g_entretien)^{A}*(1+taux_taxes))",
            f"=IF({A}=0,0,infra_entretien_pct*{INFRA}*(1+g_entretien)^{A}*(1+taux_taxes))",  # F opex
            # G résiduels alt : an 10 = reprise géométrique planchée des
            # anciens VL ; an 12 = VR de tout le monde + infra
            f"=IF({A}={ANREPL},{NV}*{VR10A},IF({A}={H_},{NV}*{VR12_VLA}+{VR12_CA}+{NM}*{VR12_MA}+{VR12_I},0))",
            f"=B{row}-C{row}+D{row}+E{row}+F{row}-G{row}",
            f"=IF({A}=0,({NV}*{PRV}+{PRC}+{NM}*{PRM})*(1+taux_taxes),"
            f"IF({A}={ANREPL},{NV}*{PRACHAT_R}*(1+taux_taxes),0))",
            f"=IF({A}=0,0,({NV}*{KMV}*{CRV}+{KMC}*{CRC}+{NM}*{KMM}*{CRM})/100*prix_diesel*(1+g_diesel)^{A}*(1+taux_taxes))",
            f"=IF({A}=0,0,({NV}*{KMV}*{ERV}+{KMC}*{ERC}+{NM}*{KMM}*{ERM})*(1+g_entretien)^{A}*(1+taux_taxes))",
            f"=IF({A}={ANREPL},{NV}*{VR10R},IF({A}={H_},{NV}*{VR12_VLR}+{VR12_CR}+{NM}*{VR12_MR},0))",
            f"=I{row}+J{row}+K{row}-L{row}",
            f"=M{row}-H{row}",
            f"=1/(1+taux_r)^{A}",
            f"=H{row}*O{row}",
            f"=M{row}*O{row}",
            (f"=N{row}" if n == 0 else f"=R{row-1}+N{row}"),
            (f"=N{row}*O{row}" if n == 0 else f"=S{row-1}+N{row}*O{row}"),
            f"=IF({A}=0,0,({NV}*{KMV}+{KMC}+{NM}*{KMM})*O{row})",
            f'=IF(R{row}>=0,{A},"")',
            f'=IF(S{row}>=0,{A},"")',
            # W TTW réf : légers/camionnette en fe_leger, camions moyens en fe_lourd
            f"=IF({A}=0,0,(({NV}*{KMV}*{CRV}+{KMC}*{CRC})*fe_leger+{NM}*{KMM}*{CRM}*fe_lourd)/100/1000)",
            f"=W{row}*ratio_wtw",
            f"=IF({A}=0,0,{KWH_TOT}*fe_elec/1000000)",
        ]
        for j, f in enumerate(formulas, start=2):
            c = ws.cell(row=row, column=j, value=f); c.fill = FILL_CALC
    last = first + H
    write_results(ws, first, last, res_row=last + 3, sorties=cas["sorties"])
    style_table(ws, hdr, last)
    return ws


# ---------------------------------------------------------------------------
def main():
    wb = Workbook()
    build_hypotheses(wb)
    par_id = {c["id"]: c for c in REF["cas"]}
    for cid in [1, 2, 3, 4, 5]:
        build_case_simple(wb, spec_depuis_json(par_id[cid]))
    build_case6(wb, par_id[6])
    build_case_simple(wb, spec_depuis_json(par_id[7]))
    out = os.path.abspath(OUT)
    wb.save(out)
    print(f"écrit : {out}")

    # vérification : les cellules de calcul contiennent des formules, pas des valeurs
    wb2 = load_workbook(out)
    n_form = 0
    for name in [f"Cas {i}" for i in range(1, 8)]:
        ws = wb2[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_form += 1
        # les cellules de résultats (colonne B du bloc Résultats) doivent être des formules
        for row in ws.iter_rows(min_col=1, max_col=2):
            lab = row[0].value
            if isinstance(lab, str) and lab.startswith(("TCO actualisé", "VAN", "Délai", "CO2e", "Coût")):
                v = row[1].value
                assert isinstance(v, str) and v.startswith("="), f"{name}!{row[1].coordinate} n'est pas une formule : {v!r}"
    print(f"vérifié : {n_form} cellules de formule ; toutes les cellules de résultats sont des formules.")


def check():
    """Évalue RÉELLEMENT les formules du classeur (bibliothèque `formulas`,
    à installer : pip install formulas) et compare aux sorties attendues du
    contre-calculateur indépendant — mêmes tolérances que les tests du
    moteur (±0,011 $ ; TCO/km ±0,00011 ; CO2 ±0,0011)."""
    import formulas  # noqa: F401 — dépendance optionnelle, hors package.json

    xl = formulas.ExcelModel().loads(os.path.abspath(OUT)).finish()
    sol = xl.calculate()
    wb = load_workbook(os.path.abspath(OUT))

    def val(sheet, cell):
        k = f"'[{os.path.basename(OUT)}]{sheet.upper()}'!{cell}"
        if k not in sol:
            return None
        v = sol[k].value
        try:
            return float(v[0, 0])
        except Exception:
            try:
                return float(v)
            except Exception:
                return str(v)

    ecarts = 0
    for cas in REF["cas"]:
        nom = f"Cas {cas['id']}"
        ws = wb[nom]
        ligne = None
        for row in ws.iter_rows(min_col=1, max_col=1):
            if isinstance(row[0].value, str) and row[0].value.startswith("TCO actualisé — alternative"):
                ligne = row[0].row
                break
        attend = cas["sorties"]
        tests = [
            (f"B{ligne}", "tcoActualiseAlt", 0.011),
            (f"B{ligne + 1}", "tcoActualiseRef", 0.011),
            (f"B{ligne + 2}", "vanDifferentielle", 0.011),
            (f"B{ligne + 4}", "tcoParKmAlt", 0.00011),
            (f"B{ligne + 8}", "co2EviteTtwTonnes", 0.0011),
            (f"B{ligne + 9}", "co2EviteWtwTonnes", 0.0011),
            (f"B{ligne + 10}", "coutParTonneWtw", 0.011),
        ]
        for cell, cle, tol in tests:
            a = attend.get(cle)
            if a is None:
                continue
            v = val(nom, cell)
            if not isinstance(v, float) or abs(v - a) > tol:
                ecarts += 1
                print(f"ÉCART {nom} {cle} : classeur={v} attendu={a}")
    if ecarts:
        raise SystemExit(f"{ecarts} écart(s) entre le classeur et le contre-calcul.")
    print("check : le classeur reproduit les 7 cas du contre-calcul au cent près.")


if __name__ == "__main__":
    import sys

    main()
    if "--check" in sys.argv:
        check()
